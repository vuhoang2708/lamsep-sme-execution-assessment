import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup - Margins (0.8 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Colors
    NAVY = RGBColor(27, 54, 93)      # #1B365D
    TEAL = RGBColor(0, 128, 128)     # #008080
    DARK_GRAY = RGBColor(60, 60, 60)
    
    # TITLE
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("BÁO CÁO RÀ SOÁT ĐỐI CHIẾU & ĐỀ XUẤT THIẾT KẾ\nBỘ CÔNG CỤ KHẢO SÁT THỰC TRẠNG THỰC THI CHO DOANH NGHIỆP SME")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY
    title_p.paragraph_format.space_after = Pt(4)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Dự án: LAMSEP - THỰC THI XUẤT SẮC (Triết lý PSO: People - Strategy - Operations)\nĐối tác thực hiện: Vũ Hoàng & Thanh Duy (DHM8) / Chuyên gia Huỳnh Trọng Nghĩa (Kenmei)")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = TEAL
    sub_p.paragraph_format.space_after = Pt(16)
    
    # Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    p_div_run = p_div.add_run("—" * 55)
    p_div_run.font.color.rgb = RGBColor(200, 200, 200)

    # ----------------------------------------------------
    # PHẦN I: BỐI CẢNH & CƠ SỞ DỮ LIỆU ĐỐI CHIẾU
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    r = h1.add_run("PHẦN I: BỐI CẢNH & CƠ SỞ DỮ LIỆU ĐỐI CHIẾU (GROUND TRUTH)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Báo cáo này được xây dựng dựa trên sự đối chiếu chéo (cross-check) toàn diện giữa hai nguồn dữ liệu thực tế:\n"
        "1. Toàn bộ nội dung trao đổi trực tiếp trong ghi âm giữa anh Vũ và anh Thanh Duy (DHM8) / Chuyên gia Huỳnh Trọng Nghĩa (Kenmei).\n"
        "2. Năm tài liệu chuyên đề chuẩn mực của bộ 'LAMSEP - Thực thi xuất sắc' bao gồm: CHIẾN LƯỢC, LÃNH ĐẠO XUẤT SẮC, VĂN HOÁ THỰC THI XUẤT SẮC, HR THỰC THI XUẤT SẮC, và VẬN HÀNH THỰC THI XUẤT SẮC."
    )
    
    p_box = doc.add_paragraph()
    p_box_run = p_box.add_run(
        "💡 Triết lý cốt lõi PSO (Larry Bossidy): Thực thi xuất sắc không phải là một công cụ riêng lẻ mà là sự liên kết nhịp nhàng giữa 3 trụ cột sống còn: Con người (People) – Chiến lược (Strategy) – Vận hành (Operations). Doanh nghiệp không vận hành bằng công việc đơn thuần mà vận hành bằng dòng chảy kết quả chuẩn hóa (Output Management)."
    )
    p_box_run.font.italic = True
    p_box_run.font.color.rgb = NAVY
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # PHẦN II: BẢNG RÀ SOÁT 5 CÂU HỎI CỐT LÕI
    # ----------------------------------------------------
    h2 = doc.add_paragraph()
    r = h2.add_run("PHẦN II: BẢNG RÀ SOÁT 5 CÂU HỎI CỐT LÕI (AUDIT MATRIX)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    # Table of 5 questions
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Câu hỏi khảo sát", "Trạng thái & Bằng chứng từ Ghi âm", "Định hướng & Quyết định chốt"]
    col_widths = [Inches(1.8), Inches(2.7), Inches(2.3)]

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1B365D")
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].width = col_widths[i]

    audit_data = [
        (
            "Câu 1: Mục đích sử dụng chính của bộ khảo sát",
            "ĐÃ CÓ THÔNG TIN RÕ RÀNG:\n- Duy nêu rõ: Muốn thiết kế một bài test/món quà tặng để doanh chủ tự nhìn ra bức tranh doanh nghiệp bị 'lủng' chỗ nào.\n- Gieo nhu cầu để khách hàng tự liên hệ tư vấn/học tập.\n- Thu thập data mặt bằng chung SME.",
            "CHỐT PHƯƠNG ÁN:\nLead Magnet kết hợp Pre-training Diagnostic Tool (Công cụ chẩn đoán thực trạng & phễu thu hút học viên/khách hàng tư vấn)."
        ),
        (
            "Câu 2: Đối tượng trực tiếp làm khảo sát",
            "ĐÃ CÓ THÔNG TIN CỤ THỂ:\n- Khối doanh nghiệp SME tại Việt Nam (chiếm 97% thị trường).\n- Duy đề cập: CEO tự đánh giá méo chỗ A, dàn quản lý làm méo chỗ B -> tạo cớ để Coach/Mentor nhảy vào đối chiếu lệch pha.",
            "CHỐT PHƯƠNG ÁN:\nChủ doanh nghiệp / CEO và Quản lý cấp trung điều hành SME (cho phép so sánh góc nhìn giữa lãnh đạo và quản lý)."
        ),
        (
            "Câu 3: Định dạng câu hỏi & Cơ chế chấm điểm",
            "ĐÃ CÓ ĐỊNH HƯỚNG TỪ DUY:\n- Duy mô phỏng giống các bài test sách quốc tế: Trả lời xong cộng điểm theo mốc (trên 80 điểm tốt, 60-70 cần cải thiện, dưới đó nguy kịch).\n- Dạng thang đo 1-5 hoặc Yes/No để tính điểm.",
            "CHỐT PHƯƠNG ÁN:\nThang đo Likert 1-5 điểm (mức độ từ Tự phát đến Chuẩn hóa xuất sắc). Tổng 50 câu = 250 điểm tối đa, quy đổi ra thang 100 điểm."
        ),
        (
            "Câu 4: Đầu ra (Output Report) mong muốn",
            "ĐÃ CÓ THÔNG TIN RẤT RÕ RÀNG:\n- Bấm nút là xuất biểu đồ mạng nhện đa giác (Radar Chart).\n- Chỉ ra doanh nghiệp bị 'méo/lủng' ở mặt nào (Thanh gỗ ngắn nhất theo nguyên lý Liebig).\n- Kèm vài dòng khuyến nghị cần cải thiện ngay.",
            "CHỐT PHƯƠNG ÁN:\nBảng điểm tổng + Biểu đồ Radar 5 trục + Báo động điểm nghẽn chí mạng + Gợi ý lộ trình hành động 30 ngày."
        ),
        (
            "Câu 5: Công cụ & Nền tảng triển khai",
            "ĐIỂM MÙ KỸ THUẬT (CẦN XÁC NHẬN):\n- Duy muốn dùng AI/tool tự động hóa nạp dữ liệu để tự vẽ biểu đồ.\n- Chưa chốt nền tảng: Web App độc lập, nhúng vào Skool/Web LàmSếp, hay Google Forms/Excel.",
            "ĐỀ XUẤT 3 LỰA CHỌN:\n1. Web App SPA tự động (Khuyên dùng - chuyên nghiệp nhất).\n2. Form online (Google Form/Tally).\n3. File Excel tự động hóa nội bộ."
        )
    ]

    for row_idx, data in enumerate(audit_data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=120, bottom=120, left=150, right=150)
            row_cells[col_idx].width = col_widths[col_idx]
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            if col_idx == 0:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # PHẦN III: BỘ 50 CÂU HỎI KHẢO SÁT CHI TIẾT
    # ----------------------------------------------------
    h3 = doc.add_paragraph()
    r = h3.add_run("PHẦN III: KHUNG 50 CÂU HỎI KHẢO SÁT THỰC TRẠNG THỰC THI (5 CHUYÊN ĐỀ)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Thang điểm đánh giá cho mỗi câu hỏi (1 đến 5 điểm):\n"
        "• 1 Điểm: Hoàn toàn chưa có / Làm hoàn toàn tự phát, cảm tính, hỗn loạn.\n"
        "• 2 Điểm: Đã có nhận thức nhưng triển khai chắp vá, không duy trì được đều đặn.\n"
        "• 3 Điểm: Đang thử nghiệm / Đã có quy trình cơ bản nhưng còn nhiều đứt gãy giữa các bộ phận.\n"
        "• 4 Điểm: Thực hiện đều đặn, có đo lường và theo dõi tiến độ rõ ràng.\n"
        "• 5 Điểm: Đã chuẩn hóa thành văn hóa tự vận hành, đo lường liên tục và tối ưu xuất sắc."
    )

    modules = [
        ("CHUYÊN ĐỀ 1: CHIẾN LƯỢC (STRATEGY) – NĂNG LỰC LỰA CHỌN & TẠO LỢI THẾ CẠNH TRANH", [
            ("Câu 1.1 (Khát vọng chiến thắng):", "Công ty có mục tiêu tăng trưởng và khát vọng chiến thắng rõ ràng (về doanh thu, thị phần, vị thế) được toàn thể đội ngũ thấu hiểu và đồng lòng hướng tới không?"),
            ("Câu 1.2 (Lựa chọn sân chơi - Where to play):", "Công ty có xác định rõ ràng phân khúc khách hàng mục tiêu, kênh bán hàng cốt lõi và kiên quyết 'từ chối' các nhóm khách hàng không phù hợp không?"),
            ("Câu 1.3 (Điểm bán hàng độc nhất - USP):", "Sản phẩm/dịch vụ của công ty có điểm khác biệt độc nhất (USP) rõ ràng giúp khách hàng nhận diện ngay so với đối thủ không?"),
            ("Câu 1.4 (Lợi thế cạnh tranh VRIN):", "Lợi thế cạnh tranh của công ty được xây dựng dựa trên nguồn lực khó sao chép, khó thay thế (mô hình VRIN) hay chỉ dựa vào việc giảm giá cạnh tranh?"),
            ("Câu 1.5 (Bẻ khóa điểm yếu đối thủ):", "Ban lãnh đạo có thường xuyên phân tích đối thủ cạnh tranh và tìm cách bẻ khóa điểm yếu của đối thủ bằng thế mạnh cốt lõi của mình không?"),
            ("Câu 1.6 (Năng lực lõi thích ứng):", "Tổ chức có năng lực thích ứng nhanh, đổi mới sản phẩm/dịch vụ khi thị trường biến động thay vì bảo thủ với mô hình cũ không?"),
            ("Câu 1.7 (Cắt giảm lãng phí & Tối ưu chi phí):", "Công ty có hệ thống nhận diện và cắt giảm triệt để các chi phí lãng phí, chi phí không tạo ra giá trị gia tăng cho khách hàng không?"),
            ("Câu 1.8 (Hệ thống đo lường chiến lược):", "Chiến lược công ty có được phân rã thành hệ thống chỉ số đo lường KPI/dữ liệu cụ thể và rà soát định kỳ hàng tháng/quý không?"),
            ("Câu 1.9 (Tránh bẫy thiên kiến xác nhận & mỏ neo):", "Khi ra quyết định kinh doanh, ban lãnh đạo có dựa trên dữ liệu thị trường thực tế thay vì bị chi phối bởi cảm tính hoặc thông tin ban đầu (Anchoring/Confirmation bias)?"),
            ("Câu 1.10 (Tránh bẫy tê liệt hành động):", "Đội ngũ lãnh đạo có tính quyết đoán cao, tránh rơi vào trạng thái 'tê liệt hành động' (Analysis Paralysis) do có quá nhiều lựa chọn không?")
        ]),
        ("CHUYÊN ĐỀ 2: LÃNH ĐẠO XUẤT SẮC (LEADERSHIP) – NGHỆ THUẬT DẪN DẮT ĐỘI NGŨ", [
            ("Câu 2.1 (Lãnh đạo làm gương - Leadership Shadow):", "Những gì lãnh đạo Nói – Làm – Ưu tiên – Đo lường có hoàn toàn nhất quán với nhau để nhân viên noi theo không?"),
            ("Câu 2.2 (Phản hồi 360 độ & Lắng nghe):", "Lãnh đạo có cởi mở tiếp nhận phản hồi 360 độ từ cấp dưới và sẵn sàng điều chỉnh hành vi của bản thân không?"),
            ("Câu 2.3 (Cân bằng Nhân trị & Pháp trị):", "Tổ chức có kết hợp hài hòa giữa đối đãi nhân văn (giữ chân nhân tài) và kỷ luật thượng tôn nguyên tắc 'quân pháp bất vị thân' không?"),
            ("Câu 2.4 (Ứng dụng Kỹ trị & Công nghệ):", "Lãnh đạo có thúc đẩy việc ứng dụng công nghệ, số hóa biểu mẫu/quy trình và dữ liệu vào điều hành hàng ngày không?"),
            ("Câu 2.5 (Thuế niềm tin - Trust Tax):", "Mức độ tin cậy nội bộ có đủ cao để các phòng ban phối hợp nhịp nhàng mà không cần quá nhiều khâu kiểm soát, phê duyệt rườm rà không?"),
            ("Câu 2.6 (Con người đi trước, công việc theo sau):", "Công ty có văn hóa tuyển chọn người phù hợp giá trị/văn hóa ngay từ đầu (Good to Great) thay vì cố chấp giữ người sai rồi tìm cách uốn nắn?"),
            ("Câu 2.7 (Sứ mệnh truyền cảm hứng):", "Lãnh đạo có truyền tải được sứ mệnh và lý tưởng công việc (Ikigai / Thuyết con nhím) giúp nhân viên tìm thấy ý nghĩa cống hiến không?"),
            ("Câu 2.8 (Dẫn dắt sự thay đổi):", "Khi cần chuyển đổi (văn hóa, công nghệ, mô hình mới), lãnh đạo có truyền thông quyết liệt và kiên trì dẫn dắt đội ngũ vượt qua lực cản không?"),
            ("Câu 2.9 (Thấu hiểu phong cách cá nhân DISC):", "Cấp quản lý có thấu hiểu phong cách hành vi của bản thân và nhân viên (DISC) để giao tiếp và phân công công việc hiệu quả không?"),
            ("Câu 2.10 (Bài binh bố trận & Quyết liệt cắt bỏ):", "Lãnh đạo có phân bổ nguồn lực tập trung vào nơi tạo doanh thu tốt nhất và quyết liệt cắt bỏ các mắt xích/dự án kém hiệu quả không?")
        ]),
        ("CHUYÊN ĐỀ 3: VĂN HOÁ THỰC THI (CULTURE) – TRẢI NGHIỆM ĐỊNH HÌNH HÀNH ĐỘNG", [
            ("Câu 3.1 (An toàn tâm lý - Psychological Safety):", "Nhân viên trong công ty có cảm thấy an toàn để nói sự thật, dám nhận lỗi và dám phản biện mà không sợ bị trừng phạt/chỉ trích không?"),
            ("Câu 3.2 (Văn hóa báo tin xấu):", "Khi có sự cố xảy ra, thông tin tiêu cực/tin xấu có được báo cáo lên cấp trên ngay lập tức thay vì bị giấu giếm không?"),
            ("Câu 3.3 (Ý nghĩa công việc - Purpose):", "Nhân viên các cấp có hiểu rõ công việc hàng ngày của họ đóng góp gì cho sự thành công của công ty và khách hàng không?"),
            ("Câu 3.4 (Ghi nhận chiến thắng nhỏ - Small Wins):", "Công ty có cơ chế ghi nhận và động viên kịp thời các bước tiến bộ nhỏ hàng tuần của nhân viên để duy trì động lực thực thi không?"),
            ("Câu 3.5 (Trách nhiệm giải trình Above-the-line):", "Khi công việc không đạt kết quả, nhân viên có chủ động nhận phần trách nhiệm (See it - Own it - Solve it - Do it) thay vì đổ lỗi hoàn cảnh?"),
            ("Câu 3.6 (Xóa bỏ tâm lý nạn nhân Below-the-line):", "Tổ chức có kiên quyết bài trừ các biểu hiện đổ lỗi: 'tại khách hàng', 'tại phòng khác', 'tại thị trường', 'tại sếp' không?"),
            ("Câu 3.7 (Xử lý 'Quả táo thối'):", "Công ty có kịp thời cô lập và xử lý các cá nhân có thái độ tiêu cực, độc hại làm lây lan và kéo tụt năng lượng của đội ngũ không?"),
            ("Câu 3.8 (Tránh hội chứng CEO anh hùng):", "Công việc có được phân quyền xử lý tự chủ hay mọi quyết định lớn nhỏ đều bị dồn về một mình CEO giải quyết?"),
            ("Câu 3.9 (Tránh bẫy hòa khí giả tạo):", "Trong các cuộc họp, mọi người có dám tranh luận thẳng thắn mang tính xây dựng hay chỉ im lặng đồng thuận bề ngoài rồi không làm?"),
            ("Câu 3.10 (Trải nghiệm định hình văn hóa):", "Công ty có chủ động kiến tạo các trải nghiệm tích cực và chính sách minh bạch để củng cố niềm tin của nhân viên không?")
        ]),
        ("CHUYÊN ĐỀ 4: HR THỰC THI (HUMAN RESOURCES) – ĐÒN BẨY NGUỒN NHÂN LỰC", [
            ("Câu 4.1 (Chiến lược thu hút nhân tài 6B):", "Công ty có chiến lược tuyển dụng rõ ràng (khi nào cần Mua ngoài - BUY, khi nào Tự đào tạo - BUILD, khi nào Thuê mượn - BORROW) không?"),
            ("Câu 4.2 (Quyết liệt đào thải - BOUNCE):", "Công ty có cơ chế sàng lọc và kiên quyết cho thôi việc/điều chuyển các nhân sự không còn phù hợp với chuẩn mực hiệu suất của tổ chức không?"),
            ("Câu 4.3 (Thương hiệu nhà tuyển dụng):", "Công ty có xây dựng được hình ảnh thương hiệu nhà tuyển dụng (Employer Branding) hấp dẫn để thu hút ứng viên chất lượng không?"),
            ("Câu 4.4 (Mô hình đào tạo 70/20/10):", "Hoạt động đào tạo có tập trung vào 70% thực chiến trên công việc và 20% kèm cặp (Mentoring/Coaching) thay vì chỉ học lý thuyết suông?"),
            ("Câu 4.5 (Đánh giá hiệu quả đào tạo Kirkpatrick):", "Các chương trình đào tạo có được đo lường dựa trên sự thay đổi hành vi (Level 3) và kết quả kinh doanh thực tế (Level 4) không?"),
            ("Câu 4.6 (Phát triển đội ngũ kế thừa):", "Công ty có lộ trình phát triển năng lực rõ ràng cho đội ngũ nhân sự chủ chốt và quản lý kế cận không?"),
            ("Câu 4.7 (Cơ cấu tổng đãi ngộ Total Rewards):", "Chính sách lương thưởng có kết hợp minh bạch giữa Lương cứng + Phụ cấp + Thưởng biến đổi theo kết quả công việc không?"),
            ("Câu 4.8 (Phúc lợi văn hóa & phi tài chính):", "Công ty có các chính sách phúc lợi mang bản sắc riêng và các hình thức động viên tinh thần (khen ngợi, tôn vinh) hiệu quả không?"),
            ("Câu 4.9 (HR theo giai đoạn phát triển):", "Chiến lược nhân sự hiện tại có bám sát đúng giai đoạn phát triển của doanh nghiệp (Startup, Tăng trưởng, Trưởng thành hay Tái cơ cấu) không?"),
            ("Câu 4.10 (Mô hình đối tác nhân sự HRBP):", "Bộ phận HR có đóng vai trò là đối tác chiến lược am hiểu kinh doanh (HRBP) hay chỉ dừng lại ở việc chấm công, tính lương sự vụ?")
        ]),
        ("CHUYÊN ĐỀ 5: VẬN HÀNH THỰC THI (OPERATIONS) – DÒNG CHẢY KẾT QUẢ TRƠN TRU", [
            ("Câu 5.1 (Nhận diện bẫy kết quả sai cách):", "Công ty có phân biệt rõ giữa thành công do may mắn/nỗ lực cơ bắp nhất thời với thành công bền vững nhờ quy trình chuẩn không?"),
            ("Câu 5.2 (Rút kinh nghiệm sau hành động - AAR):", "Sau mỗi chiến dịch hoặc dự án, đội ngũ có ngồi lại thực hiện AAR (After-Action Review) để đúc kết bài học cải tiến không?"),
            ("Câu 5.3 (Kỷ luật báo cáo Horenso):", "Nhân viên có thói quen thực hiện Horenso (Báo cáo - Liên lạc - Thảo luận) 5 phút mỗi ngày để nhận diện rủi ro từ sớm không?"),
            ("Câu 5.4 (Kế hoạch làm việc ngày & tuần):", "Các cấp quản lý và nhân viên có lập kế hoạch làm việc rõ ràng theo mức độ ưu tiên (ma trận Eisenhower) vào đầu tuần/đầu ngày không?"),
            ("Câu 5.5 (Quy trình họp tinh gọn 50 phút):", "Các cuộc họp trong công ty có tuân thủ khung giờ nghiêm ngặt, rà soát chỉ số, tập trung xử lý vấn đề và chốt kế hoạch hành động không?"),
            ("Câu 5.6 (Tránh 12 sai lầm kinh điển khi họp):", "Công ty có loại bỏ được tình trạng: họp không có mục tiêu, không có tài liệu trước, sếp nói một chiều và biên bản không có người chịu trách nhiệm?"),
            ("Câu 5.7 (Quy trình 12 bước giao việc):", "Khi giao việc, người quản lý có nói rõ kỳ vọng kết quả, yêu cầu nhân viên xác nhận lại thông tin và chốt hạn chót (Deadline) cụ thể không?"),
            ("Câu 5.8 (Quản trị theo tiêu chuẩn đầu ra - Output Management):", "Tổ chức có quản trị dựa trên chất lượng đầu ra (Output) thay vì soi xét quản lý thời gian/đầu vào (Input) của nhân viên không?"),
            ("Câu 5.9 (Chuẩn hóa điểm bàn giao giữa các phòng ban):", "Điểm bàn giao công việc giữa các phòng ban có được chuẩn hóa rõ ràng theo nguyên tắc 'Phòng tạo Output là ân nhân, phòng nhận Input là khách hàng' không?"),
            ("Câu 5.10 (Phá vỡ hiệu ứng cục bộ Silo):", "Công ty có thiết lập KPI liên đới và các nhóm dự án phối hợp chéo (Cross-functional team) để triệt tiêu tình trạng cát cứ, giấu thông tin không?")
        ])
    ]

    for mod_title, q_list in modules:
        h_mod = doc.add_paragraph()
        r_mod = h_mod.add_run(mod_title)
        r_mod.font.size = Pt(12)
        r_mod.font.bold = True
        r_mod.font.color.rgb = TEAL
        h_mod.paragraph_format.space_before = Pt(10)
        h_mod.paragraph_format.space_after = Pt(4)

        for q_code, q_text in q_list:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.2)
            p_q.paragraph_format.space_after = Pt(2)
            r_code = p_q.add_run(f"{q_code} ")
            r_code.font.bold = True
            r_code.font.color.rgb = NAVY
            p_q.add_run(q_text)

    # ----------------------------------------------------
    # PHẦN IV: CƠ CHẾ CHẤM ĐIỂM & BÁO CÁO ĐẦU RA
    # ----------------------------------------------------
    h4 = doc.add_paragraph()
    r = h4.add_run("PHẦN IV: CƠ CHẾ CHẤM ĐIỂM, PHÂN HẠNG TRƯỞNG THÀNH & BÁO CÁO RADAR")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "1. Cơ chế tính điểm:\n"
        "• Điểm từng chuyên đề: Tổng điểm 10 câu (thang điểm 10 – 50 điểm).\n"
        "• Điểm tổng thực thi toàn diện: Tổng điểm 5 chuyên đề (thang điểm 50 – 250 điểm).\n"
        "• Điểm phần trăm chuẩn hóa: (Tổng điểm đạt được / 250) x 100%."
    )

    doc.add_paragraph(
        "2. Bảng phân hạng mức độ trưởng thành thực thi của Doanh nghiệp SME:\n"
        "• Cấp độ 1: Dưới 50% (Dưới 125 điểm) – KHỞI PHÁT & HỖN LOẠN: Vận hành hoàn toàn tự phát, phụ thuộc vào cá nhân lãnh đạo, rủi ro đứt gãy rất cao.\n"
        "• Cấp độ 2: 50% - 69% (125 - 174 điểm) – THỬ NGHIỆM & CHẮP VÁ: Đã có ý thức xây dựng hệ thống nhưng quy trình chưa đồng bộ, họp nhiều làm ít, hay đổ lỗi.\n"
        "• Cấp độ 3: 70% - 84% (175 - 212 điểm) – CHUẨN HÓA & ĐỒNG BỘ: Bộ máy vận hành tương đối trơn tru, có đo lường KPI, văn hóa trách nhiệm khá tốt.\n"
        "• Cấp độ 4: 85% - 100% (213 - 250 điểm) – THỰC THI XUẤT SẮC: Văn hóa tự chủ cao (Above-the-line), dòng chảy kết quả thông suốt, tự động hóa cao."
    )

    doc.add_paragraph(
        "3. Thuật toán nhận diện 'Thanh gỗ ngắn nhất' (Liebig's Law):\n"
        "Hệ thống tự động so sánh điểm của 5 chuyên đề và chỉ ra chuyên đề có điểm thấp nhất (ví dụ: Chiến lược 42/50 nhưng Vận hành chỉ 18/50). Báo cáo sẽ đưa ra cảnh báo điểm nghẽn chí mạng này chính là nguyên nhân làm rò rỉ lợi nhuận và kéo tụt toàn bộ cỗ máy doanh nghiệp."
    )

    # ----------------------------------------------------
    # PHẦN V: ĐỀ XUẤT HÀNH ĐỘNG TIẾP THEO GỬI DUY
    # ----------------------------------------------------
    h5 = doc.add_paragraph()
    r = h5.add_run("PHẦN V: ĐỀ XUẤT HÀNH ĐỘNG TIẾP THEO (NEXT STEPS)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Để tiến hành triển khai công cụ khảo sát vào thực tế, anh Vũ có thể gửi thông điệp sau cho anh Duy để chốt bước thực thi kỹ thuật:"
    )

    p_msg = doc.add_paragraph()
    p_msg_run = p_msg.add_run(
        "\"Hi Duy, anh đã hoàn thiện bản Báo cáo rà soát và khung 50 câu hỏi khảo sát thực trạng thực thi cho DN SME dựa trên đúng 5 chuyên đề gốc của LAMSEP. Mọi định hướng về mục đích (Lead Magnet / Pre-diagnostic), đối tượng (SME CEO & Quản lý), thang điểm 1-5 và biểu đồ Radar đã được đồng bộ hoàn toàn.\n\n"
        "Hiện tại anh đề xuất dựng trước một bản Web App tương tác tự động (Single Page App) chạy mượt trên cả điện thoại và máy tính để team LàmSếp test thử nghiệm và lấy phản hồi ngay. Em xem qua file báo cáo đính kèm và cho anh ý kiến nhé!\""
    )
    p_msg_run.font.italic = True
    p_msg_run.font.color.rgb = NAVY
    p_msg.paragraph_format.left_indent = Inches(0.3)
    p_msg.paragraph_format.space_after = Pt(12)

    # Save document
    out_path = r"C:\Users\vu.hoang\.gemini\antigravity\scratch\Lam-Sep\BAO_CAO_RA_SOAT_VA_BO_KHAO_SAT_THUC_THI_SME_LAMSEP.docx"
    doc.save(out_path)
    print(f"Successfully generated docx at: {out_path}")

if __name__ == '__main__':
    create_document()
